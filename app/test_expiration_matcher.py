"""
Tests for matcher.py -- both halves: reading the expiration list, and stamping
an order confirmation / invoice without disturbing anything else.

Every fixture is synthetic and written in-memory from the literals below, so
there is no customer data here and this suite can be copied into another repo
and run as-is. The order-confirmation fixture deliberately mimics the real
file's awkward shape: column A unused, a title block above the table, headers
on row 12, and a totals block below a blank row.

Run:  python -m unittest test_expiration_matcher -v
"""
from __future__ import annotations

import datetime
import os
import tempfile
import unittest

import openpyxl

import expiration_matcher as matcher

D = datetime.datetime

# --- expiration list fixture: C&S Code | Description | UPC | BOH | Expiration Date ---
EXP_ROWS = [
    ["C&S Code", "Description", "UPC", "BOH", "Expiration Date"],
    [36, "VAGSTAT SUPP", "00-11509-06141", 44, D(2027, 7, 20)],
    [36, "VAGSTAT SUPP", "00-11509-06141", 44, D(2028, 11, 30)],   # duplicate, later
    [98, "LRL EYELINER", "00-71249-38603", 29, None],              # listed, no date
    [145, "BLU BNY FUDGE", "00-70640-01244", 295, D(2028, 2, 4)],
    [145, "BLU BNY FUDGE", "00-70640-01244", 295, D(2027, 11, 28)],  # duplicate, earlier
    ["S29756", "SILK ALMD MLK", "00-25293-00123", 12, D(2026, 12, 1)],  # letter code
]

# Header labels of the real order confirmation, starting at column B.
OC_HEADERS = [
    "Line No", "Item No", "Brand", "Pack", "Size", "Description", "COOL",
    "QTY", "Each Net WT", "Each Gross", "Unit Cube", "Unit Price", "Total Amount",
]
OC_HEADER_ROW = 12
OC_FIRST_COL = 2  # column B -- column A is empty in the real files


def write_expiration_list(path, rows=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r, row in enumerate(rows or EXP_ROWS, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    wb.save(path)
    wb.close()


def write_order_confirmation(path, item_numbers, extra_headers=(), footer=True):
    """Build an order confirmation shaped like the operator's real ones.

    Returns nothing; the caller knows the layout from the constants above.
    `extra_headers` are appended after "Total Amount" (used to test a file that
    already has its own Expiration column).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AGS_CS_Proforma_By_Order"

    # Title block above the table, as the real export has.
    ws.cell(row=2, column=OC_FIRST_COL, value="ORDER CONFIRMATION")
    ws.cell(row=3, column=OC_FIRST_COL, value="SUN COMMODITIES INC")
    ws.cell(row=3, column=13, value="Order No:")
    ws.cell(row=3, column=14, value=44678)
    ws.cell(row=7, column=OC_FIRST_COL, value="Sold To:")

    headers = list(OC_HEADERS) + list(extra_headers)
    for i, label in enumerate(headers):
        ws.cell(row=OC_HEADER_ROW, column=OC_FIRST_COL + i, value=label)

    for j, item_no in enumerate(item_numbers):
        row = OC_HEADER_ROW + 1 + j
        ws.cell(row=row, column=OC_FIRST_COL, value=j + 1)          # Line No
        ws.cell(row=row, column=OC_FIRST_COL + 1, value=item_no)    # Item No
        ws.cell(row=row, column=OC_FIRST_COL + 5, value="DESCRIPTION %d" % (j + 1))
        qty = ws.cell(row=row, column=OC_FIRST_COL + 7, value=10 + j)
        qty.number_format = "#,##0"
        price = ws.cell(row=row, column=OC_FIRST_COL + 11, value=2.5)
        total = ws.cell(row=row, column=OC_FIRST_COL + 12)
        total.value = "=%s*%s" % (qty.coordinate, price.coordinate)  # a real formula
        total.number_format = "[$-10409]#,##0.00;\\(#,##0.00\\)"

    if footer:
        # Blank row, then a totals block -- must be left alone.
        foot = OC_HEADER_ROW + 1 + len(item_numbers) + 1
        ws.cell(row=foot, column=OC_FIRST_COL + 5, value="Cubit Feet")
        ws.cell(row=foot, column=OC_FIRST_COL + 12, value="INVOICEAMT")
        ws.cell(row=foot + 1, column=OC_FIRST_COL + 5, value=17.19)
        ws.cell(row=foot + 1, column=OC_FIRST_COL + 12, value=1896.16)

    wb.save(path)
    wb.close()


class TempDir:
    def __enter__(self):
        self.dir = tempfile.mkdtemp()
        return self.dir

    def __exit__(self, *exc):
        for name in os.listdir(self.dir):
            try:
                os.unlink(os.path.join(self.dir, name))
            except OSError:
                pass
        try:
            os.rmdir(self.dir)
        except OSError:
            pass
        return False

    def path(self, name):
        return os.path.join(self.dir, name)


def item_col_letter(offset=1):
    return openpyxl.utils.get_column_letter(OC_FIRST_COL + offset)


# ===================== reading the expiration list =====================


class CleanItemKeyTests(unittest.TestCase):
    """Item numbers arrive in several shapes and must normalize to one string."""

    def test_int_becomes_plain_string(self):
        self.assertEqual(matcher._clean_item_key(97310), "97310")

    def test_whole_float_loses_its_decimal(self):
        self.assertEqual(matcher._clean_item_key(97310.0), "97310")

    def test_text_is_stripped(self):
        # The real 44676 order confirmation stores item numbers space-padded.
        self.assertEqual(matcher._clean_item_key("    107452"), "107452")

    def test_trailing_dot_zero_is_dropped(self):
        self.assertEqual(matcher._clean_item_key("97310.0"), "97310")

    def test_letters_survive(self):
        # S29756 is a real code; keys must stay strings, never become ints.
        self.assertEqual(matcher._clean_item_key("S29756"), "S29756")

    def test_blank_and_none_are_not_keys(self):
        self.assertIsNone(matcher._clean_item_key(None))
        self.assertIsNone(matcher._clean_item_key("   "))

    def test_int_and_padded_text_collide_on_purpose(self):
        self.assertEqual(matcher._clean_item_key(107452), matcher._clean_item_key("    107452"))


class LoadExpirationMapTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TempDir()
        self.dir = self.tmp.__enter__()
        self.path = self.tmp.path("exp.xlsx")
        write_expiration_list(self.path)

    def tearDown(self):
        self.tmp.__exit__()

    def test_earliest_wins_when_the_later_date_came_first(self):
        self.assertEqual(matcher.load_expiration_map(self.path)["36"], D(2027, 7, 20))

    def test_earliest_wins_when_the_earlier_date_came_second(self):
        self.assertEqual(matcher.load_expiration_map(self.path)["145"], D(2027, 11, 28))

    def test_listed_without_a_date_maps_to_none(self):
        exp_map = matcher.load_expiration_map(self.path)
        self.assertIn("98", exp_map)
        self.assertIsNone(exp_map["98"])

    def test_absent_code_is_simply_missing(self):
        self.assertNotIn("999999", matcher.load_expiration_map(self.path))

    def test_keys_are_always_strings(self):
        self.assertTrue(all(isinstance(k, str) for k in matcher.load_expiration_map(self.path)))

    def test_letter_code_is_read(self):
        self.assertEqual(matcher.load_expiration_map(self.path)["S29756"], D(2026, 12, 1))

    def test_headers_below_a_title_block(self):
        path = self.tmp.path("titled.xlsx")
        write_expiration_list(path, [["REPORT"], [], []] + EXP_ROWS)
        self.assertEqual(matcher.load_expiration_map(path)["36"], D(2027, 7, 20))

    def test_alternative_header_spellings(self):
        path = self.tmp.path("alt.xlsx")
        write_expiration_list(path, [["Item Number", "Expiration"], [1, D(2027, 1, 1)]])
        self.assertEqual(matcher.load_expiration_map(path)["1"], D(2027, 1, 1))

    def test_header_matching_ignores_case_and_padding(self):
        path = self.tmp.path("case.xlsx")
        write_expiration_list(path, [["  c&s CODE ", " EXPIRATION date "], [1, D(2027, 1, 1)]])
        self.assertEqual(matcher.load_expiration_map(path)["1"], D(2027, 1, 1))

    def test_missing_date_column_is_a_friendly_error(self):
        path = self.tmp.path("nodate.xlsx")
        write_expiration_list(path, [["C&S Code", "Description"], [1, "NOPE"]])
        with self.assertRaises(matcher.MatcherError):
            matcher.load_expiration_map(path)

    def test_headers_beyond_the_search_window_are_not_found(self):
        path = self.tmp.path("deep.xlsx")
        padding = [["FILLER"] for _ in range(matcher.MAX_HEADER_SEARCH_ROWS + 1)]
        write_expiration_list(path, padding + EXP_ROWS)
        with self.assertRaises(matcher.MatcherError):
            matcher.load_expiration_map(path)


# ===================== stamping the order confirmation =====================


class StampingTests(unittest.TestCase):
    """The half the operator's output actually depends on."""

    def setUp(self):
        self.tmp = TempDir()
        self.dir = self.tmp.__enter__()
        self.exp_path = self.tmp.path("exp.xlsx")
        write_expiration_list(self.exp_path)
        self.exp_map = matcher.load_expiration_map(self.exp_path)

    def tearDown(self):
        self.tmp.__exit__()

    def run_stamp(self, item_numbers, extra_headers=(), footer=True):
        src = self.tmp.path("oc.xlsx")
        out = self.tmp.path("out.xlsx")
        write_order_confirmation(src, item_numbers, extra_headers, footer)
        stats = matcher.process_order_confirmation(src, self.exp_map, out)
        wb = openpyxl.load_workbook(out, data_only=False)
        return stats, wb, wb.worksheets[0]

    def test_expiration_column_is_appended_when_absent(self):
        _, wb, ws = self.run_stamp([36, 145])
        # Table occupies B..N (14); the new column lands at O.
        self.assertEqual(ws.cell(row=OC_HEADER_ROW, column=15).value, "Expiration")
        wb.close()

    def test_matched_dates_are_written_as_dates_with_a_date_format(self):
        _, wb, ws = self.run_stamp([36, 145])
        self.assertEqual(ws.cell(row=13, column=15).value, D(2027, 7, 20))
        self.assertEqual(ws.cell(row=14, column=15).value, D(2027, 11, 28))
        self.assertEqual(ws.cell(row=13, column=15).number_format, "mm-dd-yy")
        wb.close()

    def test_code_absent_from_the_list_becomes_not_available(self):
        _, wb, ws = self.run_stamp([999999])
        self.assertEqual(ws.cell(row=13, column=15).value, matcher.NOT_FOUND_VALUE)
        wb.close()

    def test_code_listed_without_a_date_also_becomes_not_available(self):
        # 98 is in the list but every one of its rows has a blank date.
        _, wb, ws = self.run_stamp([98])
        self.assertEqual(ws.cell(row=13, column=15).value, matcher.NOT_FOUND_VALUE)
        wb.close()

    def test_padded_text_item_numbers_match(self):
        _, wb, ws = self.run_stamp(["        36"])
        self.assertEqual(ws.cell(row=13, column=15).value, D(2027, 7, 20))
        wb.close()

    def test_letter_code_matches(self):
        _, wb, ws = self.run_stamp(["S29756"])
        self.assertEqual(ws.cell(row=13, column=15).value, D(2026, 12, 1))
        wb.close()

    def test_stats_count_matched_and_not_found(self):
        stats, wb, _ = self.run_stamp([36, 145, 999999, 98])
        self.assertEqual(stats, {"matched": 2, "not_found": 2, "total": 4})
        wb.close()

    def test_an_existing_expiration_column_is_reused_not_duplicated(self):
        _, wb, ws = self.run_stamp([36], extra_headers=["Expiration"])
        self.assertEqual(ws.cell(row=OC_HEADER_ROW, column=15).value, "Expiration")
        self.assertIsNone(ws.cell(row=OC_HEADER_ROW, column=16).value)
        self.assertEqual(ws.cell(row=13, column=15).value, D(2027, 7, 20))
        wb.close()

    def test_reprocessing_an_already_stamped_file_is_idempotent(self):
        src = self.tmp.path("oc.xlsx")
        once = self.tmp.path("once.xlsx")
        twice = self.tmp.path("twice.xlsx")
        write_order_confirmation(src, [36, 145, 999999])
        first = matcher.process_order_confirmation(src, self.exp_map, once)
        second = matcher.process_order_confirmation(once, self.exp_map, twice)
        self.assertEqual(first, second)
        wb = openpyxl.load_workbook(twice)
        ws = wb.worksheets[0]
        self.assertEqual(ws.max_column, 15)  # no second Expiration column
        self.assertEqual(ws.cell(row=13, column=15).value, D(2027, 7, 20))
        wb.close()

    def test_the_footer_block_is_left_alone(self):
        # Items on 13-14, blank 15, footer on 16-17.
        _, wb, ws = self.run_stamp([36, 145])
        self.assertEqual(ws.cell(row=16, column=OC_FIRST_COL + 12).value, "INVOICEAMT")
        self.assertEqual(ws.cell(row=17, column=OC_FIRST_COL + 12).value, 1896.16)
        # Nothing stamped into the blank row or the footer.
        for row in (15, 16, 17):
            self.assertIsNone(ws.cell(row=row, column=15).value)
        wb.close()

    def test_stats_stop_at_the_blank_row(self):
        stats, wb, _ = self.run_stamp([36, 145])
        self.assertEqual(stats["total"], 2)  # not 4, despite the footer rows
        wb.close()

    def test_formulas_and_formats_in_other_columns_survive(self):
        _, wb, ws = self.run_stamp([36])
        total = ws.cell(row=13, column=OC_FIRST_COL + 12)
        self.assertEqual(total.value, "=I13*M13")
        self.assertIn("$-10409", total.number_format)
        self.assertEqual(ws.cell(row=13, column=OC_FIRST_COL + 7).number_format, "#,##0")
        wb.close()

    def test_the_title_block_and_sheet_name_survive(self):
        _, wb, ws = self.run_stamp([36])
        self.assertEqual(ws.title, "AGS_CS_Proforma_By_Order")
        self.assertEqual(ws.cell(row=2, column=OC_FIRST_COL).value, "ORDER CONFIRMATION")
        self.assertEqual(ws.cell(row=3, column=14).value, 44678)
        wb.close()

    def test_other_cells_on_an_item_row_are_untouched(self):
        _, wb, ws = self.run_stamp([36])
        self.assertEqual(ws.cell(row=13, column=OC_FIRST_COL).value, 1)
        self.assertEqual(ws.cell(row=13, column=OC_FIRST_COL + 5).value, "DESCRIPTION 1")
        self.assertEqual(ws.cell(row=13, column=OC_FIRST_COL + 7).value, 10)
        wb.close()

    def test_a_file_with_no_footer_still_works(self):
        stats, wb, ws = self.run_stamp([36, 145], footer=False)
        self.assertEqual(stats["total"], 2)
        self.assertEqual(ws.cell(row=14, column=15).value, D(2027, 11, 28))
        wb.close()

    def test_missing_item_no_column_is_a_friendly_error(self):
        src = self.tmp.path("bad.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="Nothing")
        ws.cell(row=2, column=1, value="useful here")
        wb.save(src)
        wb.close()
        with self.assertRaises(matcher.MatcherError):
            matcher.process_order_confirmation(src, self.exp_map, self.tmp.path("o.xlsx"))

    def test_the_source_file_is_not_modified(self):
        src = self.tmp.path("oc.xlsx")
        write_order_confirmation(src, [36])
        with open(src, "rb") as f:
            before = f.read()
        matcher.process_order_confirmation(src, self.exp_map, self.tmp.path("out.xlsx"))
        with open(src, "rb") as f:
            self.assertEqual(before, f.read())


if __name__ == "__main__":
    unittest.main(verbosity=2)


class ExpirationListAnyTabTests(unittest.TestCase):
    """The multi-distributor loader: finds the right tab by its columns, closest date wins, lists merge."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def _book(self, sheets):
        path = os.path.join(self.tmp.name, "list.xlsx")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for title, rows in sheets.items():
            ws = wb.create_sheet(title)
            for row in rows:
                ws.append(row)
        wb.save(path)
        return path

    def test_old_single_tab_layout_still_loads_with_closest_date(self):
        path = os.path.join(self.tmp.name, "old.xlsx")
        write_expiration_list(path)
        loaded, tab = matcher.load_expiration_map_any_tab(path)
        self.assertEqual(loaded["36"], D(2027, 7, 20))
        self.assertEqual(loaded["145"], D(2027, 11, 28))
        self.assertNotIn("98", loaded)  # listed without a date
        self.assertEqual(loaded["S29756"], D(2026, 12, 1))

    def test_picks_the_detail_tab_not_the_summary_tab(self):
        path = self._book({
            "Summary": [["Title"], [], ["Item #", "Cases", "Best Expiration"], ["4476", 30, D(2027, 8, 4)]],
            "Lot Detail": [["Lots"], [], ["Item #", "Expiration Date", "On Hand"],
                           ["4476", D(2027, 8, 4), 60], ["4476", D(2027, 4, 10), 28]],
            "Method": [["Notes only"]],
        })
        loaded, tab = matcher.load_expiration_map_any_tab(path)
        self.assertEqual(tab, "Lot Detail")
        self.assertEqual(loaded["4476"], D(2027, 4, 10))  # the closest, not the summary's furthest-out date

    def test_works_whichever_position_or_name_the_tab_has(self):
        path = self._book({
            "Zzz": [["Item Number", "Exp Date"], [100, D(2027, 1, 5)]],
            "Notes": [["nothing here"]],
        })
        loaded, tab = matcher.load_expiration_map_any_tab(path)
        self.assertEqual((tab, loaded), ("Zzz", {"100": D(2027, 1, 5)}))

    def test_text_dates_are_understood(self):
        path = self._book({"Sheet1": [["Item No", "Expiration Date"], ["77", "03/15/2027"], ["78", "2027-04-01"]]})
        loaded, _ = matcher.load_expiration_map_any_tab(path)
        self.assertEqual(loaded, {"77": D(2027, 3, 15), "78": D(2027, 4, 1)})

    def test_no_usable_tab_is_a_friendly_error(self):
        path = self._book({"Sheet1": [["Name", "Color"], ["a", "b"]]})
        with self.assertRaises(matcher.MatcherError):
            matcher.load_expiration_map_any_tab(path)

    def test_merge_keeps_the_closest_date_across_lists(self):
        merged = matcher.merge_expiration_maps([
            {"1": D(2027, 5, 1), "2": D(2027, 1, 1)},
            {"1": D(2027, 2, 1), "3": D(2028, 1, 1)},
            {"2": D(2027, 9, 9)},
        ])
        self.assertEqual(merged, {"1": D(2027, 2, 1), "2": D(2027, 1, 1), "3": D(2028, 1, 1)})


class ColumnLookTests(unittest.TestCase):
    """The new Expiration column must be readable on opening and look like part of the invoice's table."""

    def setUp(self):
        self.tmp = TempDir()
        self.dir = self.tmp.__enter__()
        self.exp_map = {"36": D(2027, 7, 20)}

    def tearDown(self):
        self.tmp.__exit__()

    def stamp(self, prepare=None, extra_headers=()):
        src, out = self.tmp.path("oc.xlsx"), self.tmp.path("out.xlsx")
        write_order_confirmation(src, [36, 999], extra_headers)
        if prepare:
            wb = openpyxl.load_workbook(src)
            prepare(wb.worksheets[0])
            wb.save(src)
        matcher.process_order_confirmation(src, self.exp_map, out)
        wb = openpyxl.load_workbook(out)
        return wb, wb.worksheets[0]

    @staticmethod
    def outline_table(ws):
        thin = openpyxl.styles.Side(style="thin", color="FF000000")
        border = openpyxl.styles.Border(left=thin, right=thin, top=thin, bottom=thin)
        for row in range(OC_HEADER_ROW, OC_HEADER_ROW + 3):  # header + the two item rows
            cell = ws.cell(row=row, column=14)  # "Total Amount", the last column before Expiration
            cell.border = border
            cell.alignment = openpyxl.styles.Alignment(horizontal="center")
            if row == OC_HEADER_ROW:
                cell.font = openpyxl.styles.Font(bold=True)
                cell.fill = openpyxl.styles.PatternFill("solid", fgColor="FFD9D9D9")

    def test_new_column_is_wide_enough_for_a_date(self):
        _, ws = self.stamp()
        self.assertGreaterEqual(ws.column_dimensions["O"].width, matcher.MIN_EXPIRATION_COLUMN_WIDTH)

    def test_a_narrow_or_hidden_existing_expiration_column_is_widened(self):
        def narrow(ws):
            ws.column_dimensions["O"].width = 3
            ws.column_dimensions["O"].hidden = True
        wb, ws = self.stamp(narrow, extra_headers=("Expiration",))
        self.assertGreaterEqual(ws.column_dimensions["O"].width, matcher.MIN_EXPIRATION_COLUMN_WIDTH)
        self.assertFalse(ws.column_dimensions["O"].hidden)

    def test_a_new_column_landing_in_a_template_defined_narrow_column_is_widened(self):
        """Real invoices pre-define a narrow width for the column after the table (6.7 wide: a date shows
        as ###### there) even though it has no header yet."""
        def predefined(ws):
            ws.column_dimensions["O"].width = 6.74
        _, ws = self.stamp(predefined)
        self.assertEqual(ws.cell(row=OC_HEADER_ROW, column=15).value, "Expiration")
        self.assertGreaterEqual(ws.column_dimensions["O"].width, matcher.MIN_EXPIRATION_COLUMN_WIDTH)

    def test_cells_copy_the_neighboring_columns_outline_fill_font_and_alignment(self):
        _, ws = self.stamp(self.outline_table)
        for row in range(OC_HEADER_ROW, OC_HEADER_ROW + 3):
            neighbor, mine = ws.cell(row=row, column=14), ws.cell(row=row, column=15)
            self.assertEqual(mine.border.left.style, "thin", row)
            self.assertEqual(mine.border.bottom.style, "thin", row)
            self.assertEqual(mine.alignment.horizontal, "center", row)
            self.assertEqual(mine.font.bold, neighbor.font.bold, row)
            self.assertEqual(mine.fill.fgColor.rgb, neighbor.fill.fgColor.rgb, row)

    def test_dates_keep_a_date_format_even_when_the_neighbor_is_a_number_column(self):
        _, ws = self.stamp(self.outline_table)
        self.assertEqual(ws.cell(row=OC_HEADER_ROW + 1, column=15).value, D(2027, 7, 20))
        self.assertIn("yy", ws.cell(row=OC_HEADER_ROW + 1, column=15).number_format)

    def test_footer_rows_below_the_table_get_no_styling(self):
        _, ws = self.stamp(self.outline_table)
        footer = OC_HEADER_ROW + 3 + 1  # blank row, then the totals block
        self.assertIsNone(ws.cell(row=footer, column=15).border.left.style)

    def test_a_grouped_column_width_entry_is_split_not_overlapped(self):
        def group(ws):
            ws.column_dimensions.group("B", "T", hidden=False)
            ws.column_dimensions["B"].width = 11
        wb, ws = self.stamp(group)
        widths = {k: (d.min, d.max, d.width) for k, d in ws.column_dimensions.items()}
        ranges = sorted((lo, hi) for lo, hi, _ in widths.values())
        for (lo1, hi1), (lo2, _hi2) in zip(ranges, ranges[1:]):
            self.assertLess(hi1, lo2, "overlapping <col> ranges would make Excel 'repair' the file: %s" % ranges)
        self.assertGreaterEqual(ws.column_dimensions["O"].width, matcher.MIN_EXPIRATION_COLUMN_WIDTH)
        self.assertEqual(ws.column_dimensions["B"].width, 11)  # the rest of the group keeps its width


class ItemPrefixTests(unittest.TestCase):
    def test_codes_without_the_letter_get_it_and_codes_with_it_do_not_get_it_twice(self):
        prefixed = matcher.add_item_prefix({"4476": D(2027, 4, 10), "S29756": D(2026, 12, 1), "s77": D(2027, 1, 1)})
        self.assertEqual(prefixed, {"S4476": D(2027, 4, 10), "S29756": D(2026, 12, 1), "s77": D(2027, 1, 1)})

    def test_two_spellings_of_one_code_keep_the_closest_date(self):
        prefixed = matcher.add_item_prefix({"123": D(2027, 5, 1), "S123": D(2027, 2, 1)})
        self.assertEqual(prefixed, {"S123": D(2027, 2, 1)})

    def test_an_invoice_code_typed_in_lowercase_still_matches(self):
        with TempDir() as folder:
            src, out = os.path.join(folder, "oc.xlsx"), os.path.join(folder, "out.xlsx")
            write_order_confirmation(src, ["s4476", "S4476", 4476])
            stats = matcher.process_order_confirmation(src, {"S4476": D(2027, 4, 10)}, out)
        self.assertEqual((stats["matched"], stats["not_found"]), (2, 1))  # the bare 4476 has no S, so no match
